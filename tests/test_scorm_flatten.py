"""Slides are flattened to images before a SCORM package is built."""

import io
import json
import shutil
import struct
import sys
import threading
import urllib.error
import urllib.request
import zipfile
import zlib
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1] / "evotix-scorm"
sys.path.insert(0, str(ROOT))

import evotix_scorm_gui as studio  # noqa: E402


def write_png(path: Path, width: int, height: int, rgb: tuple[int, int, int]) -> None:
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def test_password_is_a_hash_not_the_password(tmp_path, monkeypatch):
    source = (ROOT / "evotix_scorm_gui.py").read_text(encoding="utf-8")
    readme = (ROOT / "README-USB.txt").read_text(encoding="utf-8")
    assert "EvotixUSB" not in source
    assert "EvotixUSB" in readme
    assert studio.password_matches("EvotixUSB")
    assert not studio.password_matches("nope")
    monkeypatch.setattr(studio, "password_record_path", lambda: tmp_path / "studio-password.hash")
    studio.write_password("secret-pass")
    stored = (tmp_path / "studio-password.hash").read_text(encoding="utf-8")
    assert "secret-pass" not in stored
    assert stored.startswith("pbkdf2_sha256$")
    assert studio.password_matches("secret-pass")
    assert not studio.password_matches("EvotixUSB")


def test_package_uses_the_image_bytes_unchanged(tmp_path):
    slide = tmp_path / "slide.png"
    write_png(slide, 640, 360, (18, 52, 171))
    original = slide.read_bytes()
    out = tmp_path / "course.zip"
    studio.build_zip(
        [{"name": "Safety", "pages": [{"kind": "image", "path": str(slide)}], "questions": []}],
        {"title": "Safety", "primary": "#FE5000"},
        out,
    )
    with zipfile.ZipFile(out) as archive:
        names = archive.namelist()
        assert "imsmanifest.xml" in names
        assert "index.html" in names
        stored = archive.read("assets/page_000.png")
        html = archive.read("index.html").decode("utf-8")
        manifest = archive.read("imsmanifest.xml").decode("utf-8")
    assert stored == original
    assert "assets/page_000.png" in html
    assert "object-fit: contain" in html
    assert "object-fit: cover" not in html
    assert "object-fit: fill" not in html
    assert "dx < 0" in html
    assert "#FE5000" in html
    assert "#006242" in html
    assert "<schemaversion>1.2</schemaversion>" in manifest
    assert 'adlcp:scormtype="sco"' in manifest
    assert "adlcp:masteryscore>1<" in manifest
    assert "bindSwipe" in html
    assert "setPointerCapture" in html
    assert "draggable = false" in html
    assert "cmi.core.lesson_status" in html
    assert "cmi.core.session_time" in html


def test_pptx_cannot_be_packed_as_a_slide(tmp_path):
    fake = tmp_path / "deck.pptx"
    fake.write_bytes(b"not a slide image")
    with pytest.raises(studio.FlattenError):
        studio.build_zip(
            [{"name": "Deck", "pages": [{"kind": "image", "path": str(fake)}]}],
            {},
            tmp_path / "course.zip",
        )


def _deck(path: Path) -> None:
    pytest.importorskip("pptx")
    pytest.importorskip("PIL")
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    first = presentation.slides.add_slide(blank)
    fill = first.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0x12, 0x34, 0xAB)
    buffer = io.BytesIO()
    Image.new("RGB", (180, 90), (220, 20, 20)).save(buffer, format="PNG")
    buffer.seek(0)
    first.shapes.add_picture(buffer, Inches(0.4), Inches(0.4), Inches(1.6), Inches(0.8))
    second = presentation.slides.add_slide(blank)
    second.background.fill.solid()
    second.background.fill.fore_color.rgb = RGBColor(0x00, 0x62, 0x42)
    presentation.save(path)


def test_flatten_keeps_the_whole_slide(tmp_path):
    if not studio.find_soffice():
        pytest.skip("LibreOffice is required to flatten slides")
    deck = tmp_path / "deck.pptx"
    _deck(deck)
    images = studio.flatten_to_pngs(deck, tmp_path / "flat", 0)
    assert len(images) == 2
    from PIL import Image

    first = Image.open(images[0])
    second = Image.open(images[1])
    assert first.size[0] >= 1600 and first.size[1] >= 900
    assert second.size == first.size
    assert first.getpixel((8, 8))[:3] == (18, 52, 171)
    assert second.getpixel((8, 8))[:3] == (0, 98, 66)
    red = 0
    pixels = first.load()
    for y in range(0, first.height, 3):
        for x in range(0, first.width, 3):
            r, g, b = pixels[x, y][:3]
            if r > 180 and g < 80 and b < 80:
                red += 1
    assert red > 100
    source = (ROOT / "evotix_scorm_gui.py").read_text(encoding="utf-8")
    assert "import pptx" not in source
    assert "from PIL" not in source
    assert "Image.new" not in source


def _request(httpd, path, data=None, headers=None):
    url = f"http://127.0.0.1:{httpd.server_address[1]}{path}"
    request = urllib.request.Request(url, data=data, headers=headers or {})
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return response.status, dict(response.headers), response.read()
    except urllib.error.HTTPError as exc:
        return exc.code, dict(exc.headers), exc.read()


def test_studio_flattens_before_generate(tmp_path):
    if not studio.find_soffice():
        pytest.skip("LibreOffice is required to flatten slides")
    _deck(tmp_path / "deck.pptx")
    studio.reset_runtime()
    httpd = studio.make_server(0)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    try:
        status, _headers, body = _request(httpd, "/")
        assert status == 200
        assert b"Flatten slides to images" in body
        status, _headers, _body = _request(
            httpd,
            "/api/preview",
            data=b"{}",
            headers={"Content-Type": "application/json"},
        )
        assert status == 401
        status, headers, _body = _request(
            httpd,
            "/api/login",
            data=json.dumps({"password": "wrong"}).encode(),
            headers={"Content-Type": "application/json"},
        )
        assert status == 401
        status, headers, _body = _request(
            httpd,
            "/api/login",
            data=json.dumps({"password": "EvotixUSB"}).encode(),
            headers={"Content-Type": "application/json"},
        )
        assert status == 200
        cookie = headers["Set-Cookie"].split(";", 1)[0]
        boundary = "----EvotixBoundary7"
        payload = (tmp_path / "deck.pptx").read_bytes()
        preface = (
            f"--{boundary}\r\n"
            'Content-Disposition: form-data; name="file"; filename="deck.pptx"\r\n'
            "Content-Type: application/octet-stream\r\n\r\n"
        ).encode()
        ending = f"\r\n--{boundary}--\r\n".encode()
        status, _headers, body = _request(
            httpd,
            "/api/preview",
            data=preface + payload + ending,
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}", "Cookie": cookie},
        )
        assert status == 200, body
        preview = json.loads(body)
        pages = preview["modules"][0]["pages"]
        assert len(pages) == 2
        assert pages[0]["w"] >= 1600
        assert "Flattened 2 whole pages" in preview["note"]
        status, _headers, image = _request(httpd, pages[0]["url"], headers={"Cookie": cookie})
        assert status == 200
        assert image.startswith(b"\x89PNG")
        status, _headers, zipped = _request(
            httpd,
            "/api/generate",
            data=json.dumps(
                {
                    "settings": {"title": "Depot induction", "transitionType": "slide"},
                    "modules": [
                        {
                            "name": "deck.pptx",
                            "pages": [{"id": page["id"], "kind": page["kind"]} for page in pages],
                        }
                    ],
                }
            ).encode(),
            headers={"Content-Type": "application/json", "Cookie": cookie},
        )
        assert status == 200
        with zipfile.ZipFile(io.BytesIO(zipped)) as archive:
            assert archive.read("assets/page_000.png") == image
            course = archive.read("index.html").decode("utf-8")
        assert "slidein" in course
        assert "Depot induction" in course
    finally:
        httpd.shutdown()
        httpd.server_close()
        studio.reset_runtime()
